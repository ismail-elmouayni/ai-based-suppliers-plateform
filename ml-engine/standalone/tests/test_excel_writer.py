"""
Unit tests for standalone.io.excel_writer.ExcelWriter

Strategy
--------
* Use ``xlsxwriter`` to write the output (as production code does).
* Re-open the written file with ``openpyxl`` to verify sheet names, cell
  values, and chart objects — openpyxl can read xlsxwriter-produced files.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

import openpyxl
from standalone.io.excel_writer import ExcelWriter
from anomaly_detection.anomaly_flag import AnomalyFlag, Severity
from consolidation.cluster_result import ClusterMember, ConsolidationCluster
from entity_resolution.vendor_mapping import VendorMapping
from vendor_scoring.vendor_score import PerformanceBand, VendorScore

# ---------------------------------------------------------------------------
# Minimal DataFrames for writer tests
# ---------------------------------------------------------------------------

def _make_scores(n: int = 6) -> list[VendorScore]:
    bands = [PerformanceBand.GREEN, PerformanceBand.GREEN, PerformanceBand.AMBER,
             PerformanceBand.AMBER, PerformanceBand.RED, PerformanceBand.INSUFFICIENT]
    return [
        VendorScore(
            run_id=0,
            canonical_vendor_name=f"VENDOR_{i}",
            category="IT",
            composite_score=round(90.0 - i * 12, 2),
            performance_band=bands[i % len(bands)],
            saving_pct_norm=0.8,
            spend_norm=0.7,
            specialization_norm=0.9,
            raw_average_saving_percent=0.15,
            raw_total_spend=100_000.0,
            raw_specialization=1.0,
            raw_purchase_count=5,
        )
        for i in range(n)
    ]


def _make_clusters() -> list[ConsolidationCluster]:
    return [
        ConsolidationCluster(run_id=0, cluster_label=0, dominant_category="IT",
                             vendor_count=3, total_spend_at_stake=500_000.0,
                             estimated_saving_pct=0.05, estimated_saving_amount=25_000.0),
        ConsolidationCluster(run_id=0, cluster_label=1, dominant_category="Logistics",
                             vendor_count=2, total_spend_at_stake=200_000.0,
                             estimated_saving_pct=0.08, estimated_saving_amount=16_000.0),
    ]


def _make_members() -> list[ClusterMember]:
    return [
        ClusterMember(cluster_label=0, canonical_vendor_name="VENDOR_A",
                      vendor_total_spend=200_000.0, categories_supplied="IT"),
        ClusterMember(cluster_label=0, canonical_vendor_name="VENDOR_B",
                      vendor_total_spend=150_000.0, categories_supplied="IT,Logistics"),
        ClusterMember(cluster_label=1, canonical_vendor_name="VENDOR_C",
                      vendor_total_spend=120_000.0, categories_supplied="Logistics"),
    ]


def _make_flags() -> list[AnomalyFlag]:
    return [
        AnomalyFlag(
            run_id=0, source_record_id=1, po_number="PO-0001",
            canonical_vendor_name="VENDOR_A", category="IT",
            original_spend=10_000.0, spend=40_000.0,
            spend_gap=30_000.0, anomaly_score=0.95, z_score=4.1,
            severity=Severity.HIGH, reason_string="SpendGap +4.1\u03c3 above group mean",
        ),
        AnomalyFlag(
            run_id=0, source_record_id=2, po_number="PO-0002",
            canonical_vendor_name="VENDOR_B", category="IT",
            original_spend=15_000.0, spend=30_000.0,
            spend_gap=15_000.0, anomaly_score=0.70, z_score=2.3,
            severity=Severity.MEDIUM, reason_string="SpendGap +2.3\u03c3 above group mean",
        ),
    ]


def _make_mapping() -> list[VendorMapping]:
    return [
        VendorMapping(raw_vendor_name="ACME LTD",  canonical_vendor_name="ACME",
                      match_score=100.0, match_method="EXACT",        resolution_run_id=0),
        VendorMapping(raw_vendor_name="ACME CORP", canonical_vendor_name="ACME",
                      match_score=92.0,  match_method="FUZZY_WRATIO", resolution_run_id=0),
    ]


def _make_raw_df() -> pd.DataFrame:
    return pd.DataFrame([
        {"Vendor": "ACME LTD",  "Category": "IT",        "Spend": 50_000.0},
        {"Vendor": "ACME CORP", "Category": "Logistics",  "Spend": 30_000.0},
        {"Vendor": "SIEMENS",   "Category": "IT",         "Spend": 80_000.0},
    ])


# ---------------------------------------------------------------------------
# Fixture: write a complete workbook and return its path
# ---------------------------------------------------------------------------

@pytest.fixture
def written_workbook(tmp_path) -> Path:
    path = tmp_path / "output.xlsx"
    with ExcelWriter(path) as writer:
        stats = {
            "total_vendors": 5,
            "total_spend": 160_000.0,
            "anomaly_count": 2,
            "cluster_count": 2,
        }
        writer.write_summary(stats, _make_raw_df())
        writer.write_entity_resolution(_make_mapping())
        writer.write_vendor_scores(_make_scores())
        writer.write_consolidation(_make_clusters(), _make_members())
        writer.write_anomaly_flags(_make_flags())
    return path


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestExcelWriter:

    def test_output_file_created(self, written_workbook):
        """The output xlsx file is created on disk."""
        assert written_workbook.exists()
        assert written_workbook.stat().st_size > 0

    def test_all_sheets_present(self, written_workbook):
        """All 6 expected sheets are present in the workbook."""
        wb = openpyxl.load_workbook(written_workbook)
        expected = {
            "Summary", "Entity_Resolution", "Vendor_Scores",
            "Consolidation_Clusters", "Consolidation_Members", "Anomaly_Flags",
        }
        assert expected == set(wb.sheetnames)

    def test_vendor_scores_sheet_has_data_rows(self, written_workbook):
        """Vendor_Scores sheet has at least 1 data row beyond the header."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Vendor_Scores"]
        # Row 1 = header; check that row 2 has content
        assert ws.max_row >= 2
        assert ws.cell(row=2, column=1).value is not None

    def test_vendor_scores_chart_embedded(self, written_workbook):
        """Vendor_Scores sheet contains at least one chart object."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Vendor_Scores"]
        assert len(ws._charts) >= 1, "Expected at least 1 chart in Vendor_Scores"

    def test_consolidation_clusters_chart_embedded(self, written_workbook):
        """Consolidation_Clusters sheet contains at least one chart."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Consolidation_Clusters"]
        assert len(ws._charts) >= 1

    def test_anomaly_flags_chart_embedded(self, written_workbook):
        """Anomaly_Flags sheet contains at least one chart (severity pie)."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Anomaly_Flags"]
        assert len(ws._charts) >= 1

    def test_summary_sheet_has_kpi_value(self, written_workbook):
        """Summary sheet contains the total vendor count somewhere in the sheet."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Summary"]
        all_values = [ws.cell(r, c).value for r in range(1, ws.max_row + 1)
                      for c in range(1, ws.max_column + 1)]
        # The KPI value for total_vendors = 5 should appear in some cell
        assert 5 in all_values

    def test_entity_resolution_has_correct_columns(self, written_workbook):
        """Entity_Resolution header row contains the expected column names."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Entity_Resolution"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "RawVendorName" in headers
        assert "CanonicalVendorName" in headers
        assert "MatchScore" in headers

    def test_anomaly_flags_severity_values_written(self, written_workbook):
        """Severity values (HIGH, MEDIUM) appear in the Anomaly_Flags sheet."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Anomaly_Flags"]
        all_values = {ws.cell(r, c).value for r in range(1, ws.max_row + 1)
                      for c in range(1, ws.max_column + 1)}
        assert "HIGH" in all_values
        assert "MEDIUM" in all_values

    def test_context_manager_closes_file(self, tmp_path):
        """ExcelWriter closes cleanly when used as a context manager."""
        path = tmp_path / "test.xlsx"
        with ExcelWriter(path) as writer:
            writer.write_entity_resolution(_make_mapping())
        # File should be readable after context exit
        wb = openpyxl.load_workbook(path)
        assert "Entity_Resolution" in wb.sheetnames

    def test_empty_flags_df_no_chart(self, tmp_path):
        """Empty AnomalyFlags list is handled gracefully (no chart, no crash)."""
        path = tmp_path / "empty_flags.xlsx"
        with ExcelWriter(path) as writer:
            writer.write_anomaly_flags([])
        wb = openpyxl.load_workbook(path)
        assert "Anomaly_Flags" in wb.sheetnames
        ws = wb["Anomaly_Flags"]
        assert len(ws._charts) == 0

    def test_summary_chart_embedded_when_raw_data_provided(self, written_workbook):
        """Summary sheet embeds a spend-by-category chart when raw_df is non-empty."""
        wb = openpyxl.load_workbook(written_workbook)
        ws = wb["Summary"]
        assert len(ws._charts) >= 1
